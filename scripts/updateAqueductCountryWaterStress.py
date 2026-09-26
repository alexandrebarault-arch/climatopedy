"""Refresh the local WRI Aqueduct 4.0 country water-stress snapshot.

Install the optional data-maintenance dependency with:
    python -m pip install -r requirements-data.txt
"""

from __future__ import annotations

import datetime as dt
import argparse
import json
import tempfile
import urllib.request
import zipfile
from pathlib import Path

from openpyxl import load_workbook


DATA_URL = "https://aqueduct.wridata.org/Aqueduct40/aqueduct-4-0-country-rankings.zip"
OUTPUT = Path(__file__).resolve().parents[1] / "src/data/aqueductCountryWaterStress.json"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--catalog-updated", required=True, help="Verified last-updated date shown in the WRI Data Explorer (YYYY-MM-DD).")
    args = parser.parse_args()
    try:
        dt.date.fromisoformat(args.catalog_updated)
    except ValueError as error:
        raise SystemExit("--catalog-updated must be a valid YYYY-MM-DD date") from error

    with tempfile.TemporaryDirectory(prefix="climatopedy-aqueduct-") as temp_dir:
        archive_path = Path(temp_dir) / "aqueduct.zip"
        urllib.request.urlretrieve(DATA_URL, archive_path)
        with zipfile.ZipFile(archive_path) as archive:
            workbook_path = next(name for name in archive.namelist() if name.lower().endswith(".xlsx"))
            archive.extract(workbook_path, temp_dir)
            extracted = Path(temp_dir) / workbook_path
            workbook = load_workbook(extracted, read_only=True, data_only=True)
            sheet = workbook["country_future"]
            headers = next(sheet.iter_rows(values_only=True))
            index = {name: i for i, name in enumerate(headers)}
            rows = []
            try:
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if row[index["indicator_name"]] != "bws" or row[index["weight"]] != "Tot":
                        continue
                    iso3 = row[index["gid_0"]]
                    year = row[index["year"]]
                    scenario = row[index["scenario"]]
                    score = row[index["score"]]
                    category = row[index["cat"]]
                    if not isinstance(iso3, str) or len(iso3) != 3 or year not in (2030, 2050, 2080):
                        raise ValueError(f"Unexpected Aqueduct country key: {iso3!r}, {year!r}")
                    if scenario not in ("opt", "bau", "pes") or not isinstance(score, (int, float)) or not 0 <= score <= 5:
                        raise ValueError(f"Unexpected Aqueduct value for {iso3}/{year}/{scenario}: {score!r}")
                    rows.append({
                        "iso3": iso3,
                        "name": row[index["name_0"]],
                        "year": int(year),
                        "scenario": scenario,
                        "score": float(score),
                        "category": int(category),
                        "label": row[index["label"]],
                    })
            finally:
                workbook.close()

    keys = {(row["iso3"], row["year"], row["scenario"]) for row in rows}
    expected_years = {2030, 2050, 2080}
    expected_scenarios = {"opt", "bau", "pes"}
    countries = {row["iso3"] for row in rows}
    expected_keys = {(iso3, year, scenario) for iso3 in countries for year in expected_years for scenario in expected_scenarios}
    if len(keys) != len(rows) or keys != expected_keys or len(rows) < 1000:
        raise ValueError(f"Incomplete or duplicate Aqueduct snapshot: {len(rows)} rows")
    rows.sort(key=lambda row: (row["iso3"], row["year"], row["scenario"]))
    data = {
        "metadata": {
            "source": "WRI Aqueduct 4.0 Country Rankings",
            "sourceUrl": "https://datasets.wri.org/datasets/aqueduct-40-current-and-future-country-rankings",
            "methodologyUrl": "https://www.wri.org/research/aqueduct-40-updated-decision-relevant-global-water-risk-indicators",
            "catalogUpdated": args.catalog_updated,
            "downloadedAt": dt.date.today().isoformat(),
            "workbookVintage": Path(workbook_path).name,
            "indicator": "bws",
            "aggregation": "Tot (total sectoral weight); WRI demand-weighted aggregation from basin estimates",
            "years": [2030, 2050, 2080],
            "scenarioLabels": {"opt": "Optimiste (SSP1–RCP2.6)", "bau": "Tendanciel (SSP3–RCP7.0)", "pes": "Pessimiste (SSP5–RCP8.5)"},
            "limitations": [
                "Outil de priorisation national agrégé; il masque les écarts entre bassins et n’est pas une mesure directe de l’accès des ménages à l’eau potable.",
                "Aqueduct prévient que ses scores de risque global ne sont pas directement validables.",
                "Les années sont des horizons publiés et ne permettent pas d’interpolation annuelle.",
            ],
        },
        "rows": rows,
    }
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Wrote {len(rows)} rows for {len({row['iso3'] for row in rows})} countries.")


if __name__ == "__main__":
    main()
